"""Control de cifras del informe PDF frente a las salidas de R.

Uso publico: python R/auxiliares/verificar_entrega_kmo.py
La revision de PowerPoint y Word es opcional y solo usa materiales locales:
    python R/auxiliares/verificar_entrega_kmo.py --materiales-exposicion
Este control documental no recalcula KMO, MSA ni Bartlett.
"""
import argparse,csv,hashlib,json,re,zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from pypdf import PdfReader

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('--materiales-exposicion', action='store_true',
                    help='Revisar tambien los PPTX, DOCX y JSON locales; requiere python-docx.')
args=parser.parse_args()
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'resultados/kmo_bartlett_20261020'
def rows(path):return list(csv.DictReader(path.open(encoding='utf-8-sig')))
def compact(s):return re.sub(r'\s+','',s)
checks=[]
def check(name,ok):checks.append({'control':name,'cumple':bool(ok)})
pdf=ROOT/'docs/informe1_metodologico/Informe1_Metodologico.pdf'
reader=PdfReader(pdf);pages=[compact(p.extract_text()) for p in reader.pages]
for tex in sorted((OUT/'tablas_tex').glob('*.tex')):
 for line in tex.read_text(encoding='utf-8').splitlines():
  parts=line.rstrip('\\ ').split('&')
  if len(parts)>1 and re.fullmatch(r'\s*\d+(?:[.,]\d+)*\s*',parts[1]):
   needle=compact(''.join(parts))
   check(tex.name+': '+parts[0].strip(),any(needle in p for p in pages))
for value in ['0,6724059','4305,6886','901,430895','0,4137','0,3926']:
 check('Resultado vigente en PDF '+value,any(value in p for p in pages))
globales=rows(OUT/'verificacion_independiente_R.csv')
msa=rows(OUT/'verificacion_MSA_R.csv')
estado=rows(OUT/'estado_verificacion_R.csv')
check('Cinco controles globales de R dentro de tolerancia',
      len(globales)==5 and all(r['cumple'].lower()=='true' for r in globales))
check('Ocho MSA cotejados individualmente en R',
      len(msa)==8 and all(r['cumple'].lower()=='true' for r in msa))
check('Estado del cotejo vigente aprobado',
      len(estado)==1 and estado[0]['estado']=='aprobado')
check('Integridad de insumos',all(r['md5_antes']==r['md5_despues'] for r in rows(OUT/'integridad_insumos.csv')))
# El log es un artefacto local ignorado por Git. Su ausencia no acredita
# una compilacion limpia: solo impide revisar ese control adicional.
log_path=ROOT/'docs/informe1_metodologico/main.log'
if log_path.exists():
 log=log_path.read_text(encoding='utf-8',errors='replace')
 check('LaTeX sin referencias indefinidas o desbordamientos',not re.search(r'undefined|Overfull|Missing character|LaTeX Warning|Package .* Warning',log))
else:
 print('Log LaTeX no disponible: control de compilacion no realizado.')
# Los archivos de exposicion se excluyen del repositorio por decision del equipo.
# Solo se leen con la opcion explicita y no se necesitan en una clonacion publica.
sync=[]
if args.materiales_exposicion:
 from docx import Document
 definitions=json.loads((ROOT/'docs/materiales_exposicion/contenido_sincronizado_20261020.json').read_text(encoding='utf-8'))
 nslides=len(definitions)
 ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
 pptx=ROOT/'docs/Presentacion_Proyecto_Grado_CNPV_SECOP.pptx'
 with zipfile.ZipFile(pptx) as z:
  title_list=[];texts=[]
  for i in range(1,len(definitions)+1):
   root=ET.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
   text=[x.text or '' for x in root.findall('.//a:t',ns)]
   title_list.append(text[0]);texts.append(' '.join(text))
  check(f'{nslides} diapositivas',sum(bool(re.fullmatch(r'ppt/slides/slide\d+\.xml',n)) for n in z.namelist())==nslides)
  check('Sin instrucciones de exposición proyectadas',not any(re.search(r'Juan David presenta|\bTransición\b|Habla |En pantalla|guion completo',t,re.I) for t in texts))
  check('KMO y Bartlett correctos en lámina 16','0,6724059' in texts[15] and '4.305,6886' in texts[15])
  if nslides==21:
   package_text=' '.join(z.read(n).decode('utf-8',errors='ignore') for n in z.namelist() if n.endswith('.xml'))
   check('Presentación sin bloque de auditoría de IA',not re.search(r'nada[ -]sin[ -]fuente|NotebookLM|\bRAG\b|qwen|Wilson|kappa',package_text,re.I))
   for i,s in enumerate(definitions,1):
    note=ET.fromstring(z.read(f'ppt/notesSlides/notesSlide{i}.xml'))
    note_text='\n'.join(x.text or '' for x in note.findall('.//a:t',ns))
    check(f'Notas y parlamento sincronizados {i}',s['script'] in note_text)
 for three in ([True] if nslides==21 else [False,True]):
  name='Guion_Proyecto_Grado_David_Juan_Diego'+('_Juan_David' if three else '')+'.docx'
  d=Document(ROOT/'docs'/name)
  headings=[p.text for p in d.paragraphs if p.text.startswith('Diapositiva ')]
  speakers=[p.text.split('  ·')[0] for p in d.paragraphs if '  ·  ' in p.text]
  scripts=[p.text for p in d.paragraphs]
  check(f'{nslides} páginas de guion identificadas '+name,len(headings)==nslides and len(speakers)==nslides)
  for i,s in enumerate(definitions):
   good=headings[i]==f'Diapositiva {i+1:02d}  '+title_list[i] and title_list[i]==s['title']
   good=good and speakers[i]==s['speaker3' if three else 'speaker2'] and s['script'] in scripts
   sync.append({'guion':name,'diapositiva':i+1,'titulo':s['title'],'persona':speakers[i],'coincide':good})
  check('Títulos, voces y parlamentos sincronizados '+name,all(r['coincide'] for r in sync if r['guion']==name))
  check(f'Miniaturas de {nslides} láminas '+name,len(d.inline_shapes)==nslides)
  if nslides==21:
   with zipfile.ZipFile(ROOT/'docs'/name) as z:
    root=ET.fromstring(z.read('word/document.xml'))
    wns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    alltext=' '.join(x.text or '' for x in root.findall('.//w:t',wns))
    check('Guion sin bloque de auditoría de IA',not re.search(r'nada[ -]sin[ -]fuente|NotebookLM|\bRAG\b|qwen|Wilson|kappa',alltext,re.I))
    expansion=ROOT/'docs/materiales_exposicion/interpretaciones_municipales_tablas_20261020.json'
    if expansion.exists():
     data=json.loads(expansion.read_text(encoding='utf-8'))
     for number,entry in data.items():
      values=[entry[k] for k in ['lectura','municipal','proyecto','sintesis','limite']]+entry['resultados']
      check('Interpretación municipal ampliada de tabla '+number,
            f'Lectura municipal de la tabla {int(number):02d}' in alltext and all(v in alltext for v in values))

salidas=[('control_entrega_exposicion_local.csv' if args.materiales_exposicion else 'control_entrega_informe.csv',checks)]
if args.materiales_exposicion:
 salidas.append(('sincronizacion_presentacion_guiones.csv',sync))
for file,data in salidas:
 with (OUT/file).open('w',newline='',encoding='utf-8') as f:
  w=csv.DictWriter(f,fieldnames=list(data[0]));w.writeheader();w.writerows(data)
artifact_paths=[pdf]
if args.materiales_exposicion:
 artifact_paths += [pptx]+list((ROOT/'docs').glob('Guion_Proyecto_Grado_David*.docx'))
hash_name='entregables_sha256.csv' if args.materiales_exposicion else 'informe_sha256.csv'
with (OUT/hash_name).open('w',newline='',encoding='utf-8') as f:
 w=csv.DictWriter(f,fieldnames=['archivo','sha256']);w.writeheader()
 for p in artifact_paths:w.writerow({'archivo':str(p.relative_to(ROOT)).replace('\\','/'),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
print('Report pages:',len(reader.pages),'checks:',len(checks),'failures:',sum(not r['cumple'] for r in checks),'slide-script pairs:',len(sync))
for r in checks:
 if not r['cumple']:print('FAILED',r['control'])
assert all(r['cumple'] for r in checks)
